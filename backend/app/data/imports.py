import csv
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.data.models import ImportBatch
from app.data.schemas import SCHEMAS
from app.data.service import json_ready, persist_rows, validate_rows
from app.models import AuditLog

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000


def parse_file(filename: str, raw: bytes) -> tuple[list[str], list[dict]]:
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError("Choose a nonempty file no larger than 2 MiB")
    extension = PurePath(filename).suffix.lower()
    try:
        if extension == ".csv":
            reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
            table = []
            for row in reader:
                if any(cell.strip() for cell in row):
                    table.append(row)
                if len(table) > MAX_ROWS + 1:
                    raise ValueError("Split files into at most 1000 data rows")
        elif extension == ".xlsx":
            with ZipFile(io.BytesIO(raw)) as archive:
                files = archive.infolist()
                if len(files) > 2000 or sum(item.file_size for item in files) > 20 * 1024 * 1024:
                    raise ValueError("Spreadsheet's expanded contents exceed the safety limit")
                if any("vbaProject" in item.filename or "externalLinks/" in item.filename for item in files):
                    raise ValueError("Macros and external spreadsheet links are not supported")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.sheetnames) != 1:
                    raise ValueError("Upload one worksheet per import")
                sheet = workbook.active
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError("Split spreadsheets into at most 1000 data rows")
                if sheet.max_column and sheet.max_column > 50:
                    raise ValueError("At most 50 columns are supported")
                table = []
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Replace spreadsheet formulas with explicit values before importing")
                    row = [cell.value.date().isoformat() if isinstance(cell.value,datetime) else cell.value.isoformat() if isinstance(cell.value,date) else "" if cell.value is None else str(cell.value) for cell in cells]
                    if any(value.strip() for value in row):
                        table.append(row)
                    if len(table) > MAX_ROWS + 1:
                        raise ValueError("Split spreadsheets into at most 1000 data rows")
            finally:
                workbook.close()
        elif extension == ".json":
            data = json.loads(raw.decode("utf-8-sig"), parse_float=Decimal, parse_constant=lambda value: (_ for _ in ()).throw(ValueError("Non-finite JSON numbers are not supported")))
            if not isinstance(data,list) or not 1 <= len(data) <= MAX_ROWS or not all(isinstance(item,dict) for item in data):
                raise ValueError("JSON must be an array of 1–1000 flat objects")
            columns = list(dict.fromkeys(key for row in data for key in row))
            if not all(isinstance(key,str) for key in columns):
                raise ValueError("Column names must be strings")
            if any(isinstance(value,(list,dict)) for row in data for value in row.values()):
                raise ValueError("JSON import records must contain scalar fields")
            table = [columns] + [["" if row.get(key) is None else str(row[key]) for key in columns] for row in data]
        else:
            raise ValueError("Supported structured formats are CSV, XLSX, and JSON")
    except (UnicodeDecodeError, json.JSONDecodeError, BadZipFile, csv.Error) as failure:
        raise ValueError("The file could not be read in the selected format") from failure
    except Exception as failure:
        if isinstance(failure,ValueError):
            raise
        raise ValueError("The spreadsheet is malformed or unsupported") from failure
    if len(table) < 2:
        raise ValueError("Include a header and at least one data row")
    columns = [str(value).strip() for value in table[0]]
    if not 1 <= len(columns) <= 50 or any(not value or len(value)>80 for value in columns) or len(columns) != len(set(columns)):
        raise ValueError("Use 1–50 distinct, nonempty column names of at most 80 characters")
    if any(len(row) != len(columns) for row in table[1:]):
        raise ValueError("Every record must have the same number of cells as the header")
    if any(len(str(value)) > 1000 for row in table[1:] for value in row):
        raise ValueError("A cell exceeds the 1000-character limit")
    return columns, [dict(zip(columns,row)) for row in table[1:]]


def transformed(batch: ImportBatch) -> list[dict]:
    return [{field:row.get(column,"") for field,column in batch.mapping.items()} for row in batch.raw_rows]


def batch_out(batch: ImportBatch) -> dict:
    return json_ready({"id":batch.id,"business_id":batch.business_id,"resource":batch.resource,"currency":batch.currency,"filename":batch.filename,"status":batch.status,"columns":batch.columns,"row_count":len(batch.raw_rows),"preview":batch.raw_rows[:10],"mapping":batch.mapping,"transformed_preview":transformed(batch)[:10],"errors":batch.errors[:200],"error_count":len(batch.errors),"valid_count":batch.valid_count,"imported_count":batch.imported_count,"is_demo":batch.is_demo,"updated_at":batch.updated_at})


def validate_batch(db, business, batch: ImportBatch, mapping: dict[str,str]):
    if batch.status == "committed":
        raise HTTPException(409,"A committed import cannot be changed")
    if not mapping or not set(mapping) <= set(SCHEMAS[batch.resource].model_fields) or not set(mapping.values()) <= set(batch.columns) or len(set(mapping.values())) != len(mapping):
        raise HTTPException(422,"Map known business fields to distinct source columns")
    batch.mapping = mapping
    parsed, errors, _ = validate_rows(db,business,batch.resource,transformed(batch))
    batch.errors = errors
    batch.valid_count = len(batch.raw_rows)-len({error["row"] for error in errors})
    batch.is_demo = any(item.is_demo for item in parsed)
    batch.status = "validated"
    batch.currency = business.currency


def create_batch(db, business, user, resource: str, filename: str, raw: bytes):
    digest = hashlib.sha256(raw).hexdigest()
    existing = db.scalar(select(ImportBatch).where(ImportBatch.business_id == business.id,ImportBatch.resource == resource,ImportBatch.content_hash == digest))
    if existing:
        return batch_out(existing)
    try:
        columns, rows = parse_file(filename,raw)
    except ValueError as failure:
        raise HTTPException(422,str(failure))
    batch = ImportBatch(business_id=business.id,resource=resource,currency=business.currency,filename=PurePath(filename).name[:200],content_hash=digest,columns=columns,raw_rows=rows,mapping={},errors=[],valid_count=0,imported_count=0)
    db.add(batch)
    mapping = {field:field for field in SCHEMAS[resource].model_fields if field in columns}
    if mapping:
        validate_batch(db,business,batch,mapping)
    db.add(AuditLog(business_id=business.id,user_id=user.id,event="import.preview_created"))
    db.commit()
    return batch_out(batch)


def commit_batch(db, business, user, batch: ImportBatch):
    if batch.status == "committed":
        return batch_out(batch)
    if batch.status != "validated" or batch.currency != business.currency:
        raise HTTPException(409,"Validate the column mapping and current reporting currency before confirming")
    parsed, errors, context = validate_rows(db,business,batch.resource,transformed(batch))
    if errors:
        batch.errors = errors
        batch.valid_count = len(batch.raw_rows)-len({error["row"] for error in errors})
        db.commit()
        raise HTTPException(409,"Import has validation errors; no records were imported. Revalidate or download the error report.")
    batch.imported_count = persist_rows(db,business,batch.resource,parsed,context)
    batch.status = "committed"
    batch.errors = []
    db.add(AuditLog(business_id=business.id,user_id=user.id,event="import.committed",target_id=str(batch.id)))
    db.commit()
    return batch_out(batch)
